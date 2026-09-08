from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation
import pytest

from app import main
from app.db import Database
from app.ingestion import extract_document_chunks


def make_pdf_bytes(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_ref}
            )
        }
    )

    safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = DecodedStreamObject()
    stream.set_data(
        f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode("latin-1")
    )
    page[NameObject("/Contents")] = writer._add_object(stream)

    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Flood Evidence", level=1)
    document.add_paragraph("Rainfall and river-level observations support the analysis.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Methodology Review"
    textbox = slide.shapes.add_textbox(80, 120, 500, 100)
    textbox.text_frame.text = "The prototype preserves source provenance for reviewer questions."
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_pdf_extraction_preserves_page_locator():
    chunks = extract_document_chunks(
        "evidence.pdf",
        make_pdf_bytes("Flood forecasting uses rainfall evidence."),
    )
    assert chunks
    assert chunks[0].locator.startswith("page 1")
    assert "rainfall" in chunks[0].text.lower()


def test_docx_extraction_preserves_paragraph_locator():
    chunks = extract_document_chunks("evidence.docx", make_docx_bytes())
    assert chunks
    assert all(chunk.locator.startswith("paragraph ") for chunk in chunks)
    assert any("rainfall" in chunk.text.lower() for chunk in chunks)


def test_pptx_extraction_preserves_slide_locator():
    chunks = extract_document_chunks("slides.pptx", make_pptx_bytes())
    assert chunks
    assert all(chunk.locator.startswith("slide 1") for chunk in chunks)
    assert any("provenance" in chunk.text.lower() for chunk in chunks)


def test_invalid_binary_documents_fail_explicitly():
    for filename in ("broken.pdf", "broken.docx", "broken.pptx"):
        with pytest.raises(ValueError):
            extract_document_chunks(filename, b"not a valid document")


@pytest.fixture()
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    database = Database(tmp_path / "format-api.db")
    database.initialize()
    monkeypatch.setattr(main, "db", database)
    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 2 * 1024 * 1024)
    with TestClient(main.app) as test_client:
        yield test_client


def create_workspace(client: TestClient) -> str:
    response = client.post("/api/workspaces", data={"name": "Format Verification"})
    assert response.status_code == 200
    return response.json()["id"]


@pytest.mark.parametrize(
    ("filename", "content", "query", "locator_prefix", "mime_type"),
    [
        (
            "evidence.pdf",
            make_pdf_bytes("Rainfall evidence appears in this PDF."),
            "rainfall",
            "page 1",
            "application/pdf",
        ),
        (
            "evidence.docx",
            make_docx_bytes(),
            "rainfall",
            "paragraph ",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        (
            "slides.pptx",
            make_pptx_bytes(),
            "provenance",
            "slide 1",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
    ],
    ids=["pdf", "docx", "pptx"],
)
def test_format_upload_and_search_preserve_locator(
    client: TestClient,
    filename: str,
    content: bytes,
    query: str,
    locator_prefix: str,
    mime_type: str,
):
    workspace_id = create_workspace(client)

    upload = client.post(
        f"/api/workspaces/{workspace_id}/documents",
        files={"file": (filename, content, mime_type)},
    )
    assert upload.status_code == 200, upload.text
    document = upload.json()
    assert document["original_filename"] == filename
    assert document["chunk_count"] >= 1

    search = client.get(
        f"/api/workspaces/{workspace_id}/search",
        params={"q": query},
    )
    assert search.status_code == 200
    results = search.json()["results"]
    assert results
    assert results[0]["filename"] == filename
    assert results[0]["locator"].startswith(locator_prefix)
