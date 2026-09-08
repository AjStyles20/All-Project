from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import hashlib
from pathlib import Path
import re

from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from pypdf import PdfReader
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".pptx"}


class DocumentExtractionError(ValueError):
    """Raised when an accepted document cannot be safely/readably extracted."""


@dataclass(frozen=True)
class TextChunk:
    index: int
    locator: str
    text: str
    content_hash: str


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def decode_text_file(data: bytes) -> str:
    """Decode UTF-8 text input. Invalid input fails explicitly instead of guessing."""
    return data.decode("utf-8")


def chunk_text(
    text: str,
    *,
    target_chars: int = 1200,
    overlap_chars: int = 150,
) -> list[TextChunk]:
    """Create deterministic paragraph-aware chunks with bounded character overlap."""
    if target_chars <= 0:
        raise ValueError("target_chars must be greater than zero")
    if overlap_chars < 0:
        raise ValueError("overlap_chars cannot be negative")
    if overlap_chars >= target_chars:
        raise ValueError("overlap_chars must be smaller than target_chars")

    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []

    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(r"\n\s*\n+", normalized)
        if paragraph.strip()
    ]

    staged: list[tuple[int, int, str]] = []
    current: list[str] = []
    current_len = 0
    start_paragraph = 1

    for paragraph_number, paragraph in enumerate(paragraphs, start=1):
        separator_len = 2 if current else 0
        incoming_len = len(paragraph) + separator_len

        if current and current_len + incoming_len > target_chars:
            body = "\n\n".join(current).strip()
            staged.append((start_paragraph, paragraph_number - 1, body))

            previous_paragraph = current[-1]
            overlap = previous_paragraph[-overlap_chars:].strip() if overlap_chars else ""
            current = [overlap, paragraph] if overlap else [paragraph]
            current_len = sum(len(item) for item in current) + 2 * (len(current) - 1)
            start_paragraph = paragraph_number - 1 if overlap else paragraph_number
        else:
            if not current:
                start_paragraph = paragraph_number
            current.append(paragraph)
            current_len += incoming_len

    if current:
        body = "\n\n".join(current).strip()
        staged.append((start_paragraph, len(paragraphs), body))

    chunks: list[TextChunk] = []
    for index, (start, end, body) in enumerate(staged):
        locator = f"paragraph {start}" if start == end else f"paragraphs {start}-{end}"
        chunks.append(
            TextChunk(
                index=index,
                locator=locator,
                text=body,
                content_hash=sha256_text(body),
            )
        )
    return chunks


def _prefix_and_reindex(
    chunks: list[TextChunk], *, locator_prefix: str, start_index: int
) -> list[TextChunk]:
    result: list[TextChunk] = []
    for offset, chunk in enumerate(chunks):
        locator = locator_prefix
        if len(chunks) > 1:
            locator = f"{locator_prefix}; {chunk.locator}"
        result.append(
            TextChunk(
                index=start_index + offset,
                locator=locator,
                text=chunk.text,
                content_hash=chunk.content_hash,
            )
        )
    return result


def _extract_pdf(data: bytes) -> list[TextChunk]:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise DocumentExtractionError("PDF could not be opened") from exc

    if reader.is_encrypted:
        raise DocumentExtractionError("Encrypted PDF is not supported")

    chunks: list[TextChunk] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            raise DocumentExtractionError(
                f"Text extraction failed on PDF page {page_number}"
            ) from exc
        if not text:
            continue
        page_chunks = chunk_text(text)
        chunks.extend(
            _prefix_and_reindex(
                page_chunks,
                locator_prefix=f"page {page_number}",
                start_index=len(chunks),
            )
        )
    return chunks


def _extract_docx(data: bytes) -> list[TextChunk]:
    try:
        document = Document(BytesIO(data))
    except (DocxPackageNotFoundError, Exception) as exc:
        raise DocumentExtractionError("DOCX could not be opened") from exc

    chunks: list[TextChunk] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_chunks = chunk_text(text)
        chunks.extend(
            _prefix_and_reindex(
                paragraph_chunks,
                locator_prefix=f"paragraph {paragraph_number}",
                start_index=len(chunks),
            )
        )
    return chunks


def _extract_pptx(data: bytes) -> list[TextChunk]:
    try:
        presentation = Presentation(BytesIO(data))
    except (PptxPackageNotFoundError, Exception) as exc:
        raise DocumentExtractionError("PPTX could not be opened") from exc

    chunks: list[TextChunk] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = (shape.text or "").strip()
            if text:
                parts.append(text)
        if not parts:
            continue
        slide_chunks = chunk_text("\n\n".join(parts))
        chunks.extend(
            _prefix_and_reindex(
                slide_chunks,
                locator_prefix=f"slide {slide_number}",
                start_index=len(chunks),
            )
        )
    return chunks


def extract_document_chunks(filename: str, data: bytes) -> list[TextChunk]:
    """Extract supported input into indexed chunks with truthful source locators."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(f"Unsupported extension: {extension or '(none)'}")

    if extension in {".txt", ".md"}:
        try:
            text = decode_text_file(data)
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError("Text file must contain valid UTF-8") from exc
        return chunk_text(text)
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".pptx":
        return _extract_pptx(data)

    raise DocumentExtractionError(f"No extractor configured for {extension}")
